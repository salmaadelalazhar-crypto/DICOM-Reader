import ast
import io
import unittest
from pathlib import Path
from unittest.mock import patch
import numpy as np
import pydicom
from pydicom.dataset import FileDataset, FileMetaDataset
from pydicom.uid import ExplicitVRLittleEndian, SecondaryCaptureImageStorage, generate_uid
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest
import streamlit as st

ROOT = Path(__file__).resolve().parents[1]


def synthetic_dicom(name="case.dcm"):
    meta = FileMetaDataset()
    meta.TransferSyntaxUID = ExplicitVRLittleEndian
    meta.MediaStorageSOPClassUID = SecondaryCaptureImageStorage
    meta.MediaStorageSOPInstanceUID = generate_uid()
    ds = FileDataset(name, {}, file_meta=meta, preamble=b"\x00" * 128)
    ds.Rows = 16
    ds.Columns = 16
    ds.SamplesPerPixel = 1
    ds.PhotometricInterpretation = "MONOCHROME2"
    ds.BitsAllocated = 16
    ds.BitsStored = 16
    ds.HighBit = 15
    ds.PixelRepresentation = 0
    ds.Modality = "CT"
    ds.BodyPartExamined = "CHEST"
    ds.PixelData = np.arange(256, dtype=np.uint16).reshape(16, 16).tobytes()
    output = io.BytesIO()
    ds.save_as(output, enforce_file_format=True)
    return output.getvalue()


class Upload(io.BytesIO):
    def __init__(self, data, name):
        super().__init__(data)
        self.name = name


class IntegrationTests(unittest.TestCase):
    def run_app(self, upload, service):
        original_uploader = st.file_uploader
        def uploader(label, *args, **kwargs):
            if label.startswith("ارفع ملف أشعة"):
                return upload
            return original_uploader(label, *args, **kwargs)
        self.uploader = uploader
        with patch("streamlit.file_uploader", side_effect=uploader):
            return AppTest.from_file(str(ROOT / "app.py"), default_timeout=60).run()

    def test_real_monai_dicom_to_gemini_then_audio_and_notes(self):
        with patch("monai_gemini.ui.GeminiService") as cls, patch("gtts.gTTS.write_to_fp", side_effect=lambda fp: fp.write(b"original-monai-mp3")):
            service = cls.return_value
            service.available_models.return_value = {"gemini-3.5-flash-lite"}
            service.resolve_model.return_value = "gemini-3.5-flash-lite"
            service.analyze.side_effect = ["Gemini first", "Gemini notes"]
            service.speech_mp3.return_value = b"independent-ai-mp3"
            app = self.run_app(Upload(synthetic_dicom(), "case.dcm"), service)
            self.assertEqual(len(app.exception), 0)
            self.assertEqual(service.analyze.call_count, 1)
            report = service.analyze.call_args.args[0]
            self.assertIn("MONAI + PyTorch", report)
            self.assertIn('"centroid"', report)
            self.assertIn('"p95"', report)
            self.assertNotIn('"norm"', report)
            with patch("streamlit.file_uploader", side_effect=self.uploader):
                app.button(key="play_0").click().run()
                self.assertEqual(len(app.get("audio")), 1)
                app.text_area(key="mg_notes").input("ملاحظة صوتية معدلة")
                app.button(key="mg_send").click().run()
                self.assertEqual(service.analyze.call_args.args[2], "ملاحظة صوتية معدلة")
                app.button(key="mg_tts").click().run()
                self.assertEqual(len(app.get("audio")), 2)
                self.assertEqual(len(app.exception), 0)
                self.assertEqual(service.analyze.call_count, 2)
                self.assertEqual(app.session_state["mg_ai_audio"], b"independent-ai-mp3")

    def test_batch_zip_selection_preserves_original_results(self):
        import zipfile
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w") as z:
            z.writestr("first.dcm", synthetic_dicom())
            z.writestr("second.dcm", synthetic_dicom())
        with patch("monai_gemini.ui.GeminiService") as cls:
            service = cls.return_value
            service.available_models.return_value = {"gemini-3.5-flash-lite"}
            service.resolve_model.return_value = "gemini-3.5-flash-lite"
            service.analyze.return_value = "تحليل مساعد"
            app = self.run_app(Upload(buffer.getvalue(), "batch.zip"), service)
            self.assertEqual(len(app.exception), 0)
            self.assertIn("first.dcm", service.analyze.call_args.args[0])
            with patch("streamlit.file_uploader", side_effect=self.uploader):
                app.selectbox(key="gemini_source_result").select(1).run()
            self.assertIn("second.dcm", service.analyze.call_args.args[0])
            self.assertEqual(service.analyze.call_count, 2)
            self.assertEqual(len(app.exception), 0)

    def test_no_secret_does_not_break_original_app(self):
        from monai_gemini.service import ServiceError
        with patch("monai_gemini.ui.GeminiService", side_effect=ServiceError("لم يضبط المفتاح")):
            app = self.run_app(Upload(synthetic_dicom(), "case.dcm"), None)
            self.assertEqual(len(app.exception), 0)
            self.assertTrue(any(x.value == "لم يضبط المفتاح" for x in app.error))
            self.assertGreaterEqual(len(app.get("download_button")), 1)

    def test_original_excel_outputs_with_real_monai(self):
        # Load only original processing definitions, never execute application UI.
        source = ast.parse((ROOT / "app.py").read_text())
        wanted = {"safe", "read_dicom", "to_2d", "monai_analysis", "level", "build_report", "build_excel", "analyze_one", "build_batch_excel", "extract_dicoms_from_zip"}
        nodes = [n for n in source.body if isinstance(n, ast.FunctionDef) and n.name in wanted]
        import pandas as pd
        import torch
        from monai.transforms import Compose, EnsureChannelFirst, ScaleIntensity, ToTensor
        scope = dict(io=io, np=np, pd=pd, pydicom=pydicom, torch=torch, Compose=Compose, EnsureChannelFirst=EnsureChannelFirst, ScaleIntensity=ScaleIntensity, ToTensor=ToTensor, MONAI_OK=True, MODALITY_AR={"CT": "أشعة مقطعية (CT)"})
        exec(compile(ast.Module(body=nodes, type_ignores=[]), "processing", "exec"), scope)
        result = scope["analyze_one"]("case.dcm", synthetic_dicom())
        self.assertEqual(result[1]["engine"], "MONAI + PyTorch")
        self.assertEqual(result[1]["norm"].shape, (16, 16))
        book = load_workbook(io.BytesIO(scope["build_excel"](*result)))
        self.assertEqual(book.sheetnames, ["البيانات التقنية", "التحليل الهيكلي", "التقرير"])
        batch = load_workbook(io.BytesIO(scope["build_batch_excel"]([result, result])))
        self.assertEqual(batch["ملخص الملفات"].max_row, 3)
        scope["MONAI_OK"] = False
        fallback = scope["analyze_one"]("case.dcm", synthetic_dicom())
        self.assertIn("NumPy", fallback[1]["engine"])
        self.assertAlmostEqual(result[1]["mean"], fallback[1]["mean"], places=6)


if __name__ == "__main__":
    unittest.main()
